"""Preserve approved source workbooks. No Tally layout is invented here.

Actual Tally sources have not been available in this environment. A registry must
be completed AFTER inspecting official workbooks and mapping XML. The tests use
synthetic fixtures solely to exercise this adapter, not certify Tally imports.
"""
from copy import copy
import hashlib
import json
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook


class TemplateError(ValueError):
    pass


def inspect_workbook(path):
    book = load_workbook(path, data_only=False)
    try:
        return {'file':Path(path).name, 'sha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),
                'sheets': {sheet.title: {'rows':sheet.max_row, 'columns':sheet.max_column,
                         'sample_rows': [[cell.value for cell in row] for row in sheet.iter_rows(max_row=min(sheet.max_row, 12))]}
                         for sheet in book}}
    finally:
        book.close()


def registry(root):
    path = root / 'registry.json'
    if not path.is_file():
        raise TemplateError('Official Tally workbooks and verified registry.json are missing; import export is blocked')
    data = json.loads(path.read_text(encoding='utf-8'))
    required = {'ledger','unit','stock_item','sales'}
    if not required.issubset(data):
        raise TemplateError('Registry must include ledger, unit, stock_item and sales official mappings')
    return data


def check_spec(root, spec):
    if spec.get('verified') is not True or not spec.get('reviewed_by'):
        raise TemplateError('Official source and its mapping must be inspected and explicitly verified')
    host = urlparse(spec.get('source_url', '')).hostname or ''
    if not (host == 'tallysolutions.com' or host.endswith('.tallysolutions.com')):
        raise TemplateError('Template provenance must identify an official Tally Solutions download URL')
    source = (root / spec['file']).resolve()
    if not source.is_relative_to(root.resolve()) or not source.is_file():
        raise TemplateError('Source workbook must be inside the shared templates folder')
    if hashlib.sha256(source.read_bytes()).hexdigest() != spec.get('sha256'):
        raise TemplateError('Source workbook digest does not match the inspected registry')
    mapping_file = (root / spec['mapping_file']).resolve()
    if not mapping_file.is_relative_to(root.resolve()) or not mapping_file.is_file():
        raise TemplateError('Official companion mapping file is missing')
    if hashlib.sha256(mapping_file.read_bytes()).hexdigest() != spec.get('mapping_sha256'):
        raise TemplateError('Official companion mapping digest mismatch')
    book = load_workbook(source)
    try:
        for block in spec['blocks']:
            sheet = book[block['sheet']]
            headers = [sheet.cell(block['header_row'], col).value for col in range(1, sheet.max_column + 1)]
            if headers != block['headers']:
                raise TemplateError(f'Exact official headers changed in {sheet.title}')
            if block['data_start_row'] <= block['header_row']:
                raise TemplateError('Data rows must follow the headers')
            if not isinstance(block.get('sample_end_row'), int) or block['sample_end_row'] < block['data_start_row']:
                raise TemplateError('Explicit inspected sample_end_row is required')
            mapped = set(block['columns'])
            mandatory = set(block['mandatory_headers'])
            if not mapped.issubset(set(headers)) or not mandatory.issubset(mapped):
                raise TemplateError('Missing mandatory official columns or unknown column mapping')
        return source, mapping_file
    finally:
        book.close()


def resolve_field(record, field):
    value = record
    for part in field.split('.'):
        if not isinstance(value, dict) or part not in value:
            raise TemplateError(f'Mapping references absent field: {field}')
        value = value[part]
    return value


def write_workbook(root, spec, output, streams):
    source, mapping = check_spec(root, spec)
    book = load_workbook(source)
    try:
        for block in spec['blocks']:
            sheet = book[block['sheet']]
            start = block['data_start_row']
            styles = [copy(sheet.cell(start, col)._style) for col in range(1, sheet.max_column + 1)]
            # Only the explicitly inspected sample data range is cleared. ReadMe,
            # sheet names, headers and other fixed cells remain intact.
            for row in sheet.iter_rows(min_row=start, max_row=block['sample_end_row']):
                for cell in row:
                    cell.value = None
            records = streams.get(block['stream'], [])
            for index, record in enumerate(records, start):
                if index > block['sample_end_row'] and any(sheet.cell(index, c).value is not None for c in range(1, sheet.max_column+1)):
                    raise TemplateError('Output would overwrite fixed template content')
                for header, rule in block['columns'].items():
                    if isinstance(rule, str):
                        value = resolve_field(record, rule)
                    elif isinstance(rule, dict) and set(rule) == {'constant'}:
                        value = rule['constant']
                    else:
                        raise TemplateError(f'Unsupported mapping rule for {header}')
                    if header in block['mandatory_headers'] and value in ('', None):
                        raise TemplateError(f'Missing mandatory official field {header}')
                    col = block['headers'].index(header) + 1
                    cell = sheet.cell(index, col)
                    cell._style = copy(styles[col - 1])
                    cell.value = value
                    # Invoice data must never become an Excel formula.
                    if isinstance(value, str):
                        cell.data_type = 's'
        output.parent.mkdir(parents=True, exist_ok=True)
        temp = output.with_suffix('.tmp.xlsx')
        book.save(temp)
        temp.replace(output)
        return mapping
    finally:
        book.close()
