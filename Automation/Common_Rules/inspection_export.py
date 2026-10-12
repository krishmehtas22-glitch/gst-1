"""Write every scanned invoice and item to a reviewable Excel workbook."""
from pathlib import Path
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from .workflow_status import invoice_status

HEADERS=(
    'Client folder','Source PDF','PDF page','Invoice number','Invoice date','Seller','Seller GSTIN',
    'Seller address','Buyer','Buyer GSTIN','Buyer registration','Buyer address','Buyer country','Seller state','Place of supply','Supply','Classification','Export mode',
    'Currency','Foreign currency','Exchange rate','Item number','Item description','HSN/SAC','Quantity','Unit','Rate','Taxable value','GST rate %','CGST rate %',
    'CGST amount','SGST rate %','SGST amount','IGST rate %','IGST amount','Invoice taxable total',
    'Invoice CGST','Invoice SGST','Invoice IGST','Invoice total','Source warnings','Processing errors',
    'Invoice status','Source PDF path','PDF SHA256')


def _cell(value):
    if value is None:return ''
    if isinstance(value,(list,tuple,set)):value='; '.join(map(str,value))
    if isinstance(value,dict):value=str(value)
    value=str(value)
    if value[:1] in ('=','+','-','@','\t','\r'):return "'"+value
    return value


def export_inspection(root,records):
    """Keep prior client rows and refresh the PDFs included in this scan."""
    root=Path(root);destination=root/'Invoice_Inspection.xlsx';existing={}
    if destination.is_file():
        try:
            old=load_workbook(destination,read_only=True,data_only=True)
            sheet=old['Invoices'] if 'Invoices' in old.sheetnames else None
            if sheet:
                iterator=sheet.iter_rows(values_only=True);old_headers=next(iterator,())
                indexes={name:i for i,name in enumerate(old_headers)}
                for values in iterator:
                    get=lambda name: values[indexes[name]] if name in indexes and indexes[name]<len(values) else ''
                    key=(str(get('Source PDF path')),str(get('PDF SHA256')),str(get('PDF page')),str(get('Invoice number')),str(get('Item number')))
                    if key[0]:existing[key]=[get(name) for name in HEADERS]
            old.close()
        except (OSError,ValueError,KeyError,TypeError):
            existing={}
    refreshed_paths={str(Path(r.get('source_path') or r.get('path') or '').resolve()) for r in records}
    existing={key:row for key,row in existing.items() if key[0] not in refreshed_paths}
    for record in records:
        invoice=record.get('invoice') or {};source=record.get('source_path') or record.get('path') or ''
        source_path=str(Path(source).resolve()) if source else ''
        common={
            'Client folder':record.get('client',''),'Source PDF':Path(source).name if source else '',
            'PDF page':record.get('page_number',''),'Invoice number':invoice.get('number',''),
            'Invoice date':invoice.get('date',''),'Seller':invoice.get('seller_name',''),
            'Seller GSTIN':invoice.get('seller_gstin',''),'Buyer':invoice.get('buyer_name',''),
            'Buyer GSTIN':invoice.get('buyer_gstin',''),'Buyer registration':invoice.get('buyer_registration',''),
            'Seller address':invoice.get('seller_address',''),'Buyer address':invoice.get('buyer_address',''),
            'Buyer country':invoice.get('buyer_country',''),'Export mode':invoice.get('export_mode',''),
            'Foreign currency':invoice.get('foreign_currency',''),'Exchange rate':invoice.get('exchange_rate',''),
            'Seller state':invoice.get('seller_state',''),'Place of supply':invoice.get('place_of_supply',''),
            'Supply':invoice.get('supply',''),'Classification':invoice.get('classification',''),
            'Currency':invoice.get('currency','INR'),'Invoice taxable total':invoice.get('taxable',''),
            'Invoice CGST':invoice.get('cgst',''),'Invoice SGST':invoice.get('sgst',''),
            'Invoice IGST':invoice.get('igst',''),'Invoice total':invoice.get('total',''),
            'Source warnings':' | '.join(map(str,invoice.get('extraction_warnings',[]))),
            'Processing errors':' | '.join(map(str,record.get('errors',[]))),
            'Invoice status':record.get('display_status') or invoice_status(record),
            'Source PDF path':source_path,'PDF SHA256':record.get('digest','')}
        items=invoice.get('items') or [{}]
        for item_index,item in enumerate(items,1):
            row=dict(common)
            row.update({'Item number':item_index if item.get('name') else '', 'Item description':item.get('name',''),'HSN/SAC':item.get('hsn',''),
                'Quantity':item.get('quantity',''),'Unit':item.get('unit',''),'Rate':item.get('rate',''),
                'Taxable value':item.get('value',''),'GST rate %':item.get('gst_rate',''),'CGST rate %':item.get('cgst_rate',''),
                'CGST amount':item.get('printed_cgst',item.get('cgst','')),
                'SGST rate %':item.get('sgst_rate',''),'SGST amount':item.get('printed_sgst',item.get('sgst','')),
                'IGST rate %':item.get('igst_rate',''),'IGST amount':item.get('printed_igst',item.get('igst',''))})
            key=(source_path,str(record.get('digest','')),str(record.get('page_number','')),
                 str(invoice.get('number','')),str(item_index if item.get('name') else ''))
            existing[key]=[_cell(row.get(name,'')) for name in HEADERS]
    book=Workbook();sheet=book.active;sheet.title='Invoices';sheet.append(HEADERS)
    sheet.freeze_panes='A2';sheet.auto_filter.ref=f'A1:{get_column_letter(len(HEADERS))}{len(existing)+1}'
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78')
    for row in existing.values():sheet.append(row)
    widths=(18,31,11,19,14,32,21,31,28,21,20,31,31,20,15,20,14,17,17,20,18,12,34,14,13,13,16,18,14,16,14,16,14,16,20,16,16,16,19,50,50,36,60,66)
    for index,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(index)].width=width
    reports=root/'reports';reports.mkdir(parents=True,exist_ok=True)
    temp=destination.with_name(destination.stem+'.tmp.xlsx')
    book.save(temp);os.replace(temp,destination)
    return destination
