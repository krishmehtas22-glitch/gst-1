from copy import deepcopy
import hashlib
import shutil
import sqlite3
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
import pytest
from openpyxl import load_workbook

from Common_Rules.accounting import validate
from Common_Rules.engine import INBOX, confirm_import, read_json, run, write_json
from Common_Rules.history import History
from Common_Rules.templates import TemplateError, check_spec, inspect_workbook
from helpers import client, fake_templates, invoice, pdf, write_snapshot


def test_25_pdfs_incremental_and_same_number_per_client(tmp_path):
    a = client(tmp_path)
    b = client(tmp_path, 'Client 100')
    for n in range(25):
        pdf(a / INBOX / f'{n:02}.pdf',invoice(f'INV-{n}', ['intra','inter','exempt'][n % 3]))
    pdf(b / INBOX / 'one.pdf',invoice('INV-0'))
    first = run(tmp_path)
    assert first['clients'][0]['counts']['new'] == 25
    assert first['clients'][0]['pending_validated'] == 25
    assert first['clients'][1]['counts']['new'] == 1
    assert first['clients'][0]['export'].startswith('Generated 25')
    assert len(ET.parse(a / 'imports' / 'Sales_Vouchers.xml').findall('.//VOUCHER')) == 25
    assert not (a / 'imports' / 'Sales_Output.xlsx').exists()
    second = run(tmp_path)
    assert second['clients'][0]['counts']['new'] == 0
    assert second['clients'][0]['counts']['duplicate'] == 25
    assert len(second['clients'][0]['history']) == 25
    for n in range(25,50):
        pdf(a / INBOX / f'{n:02}.pdf',invoice(f'INV-{n}'))
    third = run(tmp_path)
    assert third['clients'][0]['counts']['new'] == 25
    assert third['clients'][0]['pending_validated'] == 50
    assert len(ET.parse(a / 'imports' / 'Sales_Vouchers.xml').findall('.//VOUCHER')) == 50
    assert len(third['clients'][1]['history']) == 1


def test_add_rename_delete_clients_and_move_source(tmp_path):
    a = client(tmp_path)
    b = client(tmp_path,'Client 2')
    pdf(a / INBOX / 'a.pdf',invoice())
    pdf(b / INBOX / 'a.pdf',invoice())
    run(tmp_path)
    nested = b / INBOX / 'moved'
    nested.mkdir()
    (b / INBOX / 'a.pdf').rename(nested / 'renamed.pdf')
    b.rename(tmp_path / 'Renamed Business')
    shutil.rmtree(a)
    c = client(tmp_path,'Client 5')
    pdf(c / INBOX / 'new.pdf',invoice('NEW'))
    result = run(tmp_path)
    assert [r['client'] for r in result['clients']] == ['Client 5','Renamed Business']
    assert result['clients'][0]['counts']['new'] == 1
    assert result['clients'][1]['counts']['duplicate'] == 1
    shutil.rmtree(c)
    shutil.rmtree(tmp_path / 'Renamed Business')
    assert run(tmp_path)['message'] == 'No client folders found'


def test_duplicate_different_content_flagged(tmp_path):
    c = client(tmp_path)
    fake_templates(tmp_path)
    pdf(c / INBOX / 'a.pdf',invoice())
    run(tmp_path)
    different = invoice()
    different['buyer_name'] = 'Different buyer'
    pdf(c / INBOX / 'b.pdf',different)
    report = run(tmp_path)['clients'][0]
    assert report['counts']['review'] == 1
    assert any('transaction details differ' in str(i) for i in report['issues'])
    assert len(report['history']) == 1
    assert report['pending_validated'] == 0
    assert not (c / 'imports' / 'Sales_Output.xlsx').exists()


def test_pdf_errors_do_not_affect_other_client(tmp_path):
    a = client(tmp_path)
    b = client(tmp_path,'Client 2')
    (a / INBOX / 'bad.pdf').write_bytes(b'not a PDF')
    pdf(b / INBOX / 'good.pdf',invoice())
    result = run(tmp_path)
    assert result['clients'][0]['counts']['review'] == 1
    assert result['clients'][1]['counts']['new'] == 1


def test_review_override_is_digest_bound_and_still_validated(tmp_path):
    c = client(tmp_path)
    source = c / INBOX / 'scanned.pdf'
    out = __import__('reportlab.pdfgen.canvas',fromlist=['Canvas']).Canvas(str(source))
    out.showPage()
    out.save()
    first = run(tmp_path)['clients'][0]
    assert first['counts']['review'] == 1
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    review = c / 'reviews' / f'{digest}.json'
    write_json(review,{'source_sha256':'wrong','approved':True,'invoice':invoice()})
    assert run(tmp_path)['clients'][0]['counts']['review'] == 1
    write_json(review,{'source_sha256':digest,'approved':True,'invoice':dict(invoice(),total='999')})
    assert run(tmp_path)['clients'][0]['counts']['review'] == 1
    write_json(review,{'source_sha256':digest,'approved':True,'invoice':invoice()})
    assert run(tmp_path)['clients'][0]['counts']['new'] == 1


def test_existing_company_requires_snapshot_and_reuses_voucher(tmp_path):
    c = client(tmp_path,mode='existing')
    pdf(c / INBOX / 'one.pdf',invoice())
    assert 'requires a complete' in run(tmp_path)['clients'][0]['error']
    config = read_json(c / 'client.json')
    write_snapshot(c / 'tally_snapshot.json',config,[{'invoice':invoice()}])
    report = run(tmp_path)['clients'][0]
    assert report['counts']['already_in_tally'] == 1
    assert report['history'][0]['status'] == 'imported'
    assert report['pending_validated'] == 0


def test_tally_conflict_and_wrong_company(tmp_path):
    c = client(tmp_path,mode='existing')
    config = read_json(c / 'client.json')
    pdf(c / INBOX / 'one.pdf',invoice())
    different = invoice()
    different['buyer_name'] = 'Existing other buyer'
    write_snapshot(c / 'tally_snapshot.json',config,[{'invoice':different}])
    assert run(tmp_path)['clients'][0]['counts']['review'] == 1
    snapshot = read_json(c / 'tally_snapshot.json')
    snapshot['company_id'] = 'wrong'
    write_json(c / 'tally_snapshot.json',snapshot)
    assert 'different client/company' in run(tmp_path)['clients'][0]['error']


def test_stale_snapshot_rejected(tmp_path):
    c = client(tmp_path,mode='existing')
    config = read_json(c / 'client.json')
    write_snapshot(c / 'tally_snapshot.json',config)
    snapshot = read_json(c / 'tally_snapshot.json')
    snapshot['exported_at'] = '2020-01-01T00:00:00+00:00'
    write_json(c / 'tally_snapshot.json',snapshot)
    assert 'stale' in run(tmp_path)['clients'][0]['error']


def test_multi_item_single_invoice(tmp_path):
    c = client(tmp_path)
    inv = invoice()
    inv['items'][0].update(quantity='1',value='1000')
    second = deepcopy(inv['items'][0])
    second['name'] = 'Product B'
    inv['items'].append(second)
    pdf(c / INBOX / 'multi.pdf',inv)
    report = run(tmp_path)['clients'][0]
    assert report['counts']['new'] == 1
    records = read_json(c / 'reports' / 'validated_invoices.json')
    assert len(records) == 1
    assert len(records[0]['invoice']['items']) == 2


def test_synthetic_adapter_preserves_headers_readme_and_history_status(tmp_path):
    c = client(tmp_path)
    fake_templates(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice())
    report = run(tmp_path)['clients'][0]
    assert report['history'][0]['status'] == 'generated'
    assert report['counts']['new'] == 1
    manifest = read_json(c / 'imports' / 'manifest.json')
    assert manifest['valid_for_import'] is True
    book = load_workbook(c / 'imports' / 'Sales_Output.xlsx')
    assert book.sheetnames == ['Synthetic Data','ReadMe']
    assert book['ReadMe']['A1'].value.startswith('SYNTHETIC')
    assert book['Synthetic Data']['A1'].value == 'Test Identifier'
    assert book['Synthetic Data']['A2'].value == 'INV-001'
    assert book['Synthetic Data']['A3'].value is None
    book.close()
    again = run(tmp_path)['clients'][0]
    assert again['counts']['new'] == 0
    assert again['counts']['duplicate'] == 1


def test_import_confirmation_requires_master_success(tmp_path):
    c = client(tmp_path)
    fake_templates(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice())
    run(tmp_path)
    manifest = read_json(c / 'imports' / 'manifest.json')
    result = {'company_id':manifest['company_id'],'generated_at':manifest['generated_at'],
              'tally_result_reference':'SYNTHETIC TEST ONLY - not actual Tally import',
              'vouchers':[{'invoice_key':manifest['invoice_keys'][0],'status':'imported'}]}
    result_path = c / 'import_results.json'
    write_json(result_path,result)
    with pytest.raises(ValueError,match='dependent masters'):
        confirm_import(c,result_path)
    result['masters'] = [{'master':m,'status':'imported'} for m in manifest['masters']]
    write_json(result_path,result)
    assert confirm_import(c,result_path)[0]['status'] == 'imported'
    report = run(tmp_path)['clients'][0]
    assert report['pending_validated'] == 0
    assert not (c / 'imports' / 'Sales_Output.xlsx').exists()


def test_template_digest_and_header_guards(tmp_path):
    folder, specs = fake_templates(tmp_path)
    spec = deepcopy(specs['sales'])
    check_spec(folder,spec)
    spec['verified'] = False
    with pytest.raises(TemplateError,match='inspected'):
        check_spec(folder,spec)
    spec = deepcopy(specs['sales'])
    spec['blocks'][0]['headers'][0] = 'Invented'
    with pytest.raises(TemplateError,match='headers changed'):
        check_spec(folder,spec)
    (folder / 'sales.xlsx').write_bytes(b'changed')
    with pytest.raises(TemplateError,match='digest'):
        check_spec(folder,specs['sales'])


def test_formula_injection_stays_text(tmp_path):
    c = client(tmp_path)
    fake_templates(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice('=1+1'))
    run(tmp_path)
    book = load_workbook(c / 'imports' / 'Sales_Output.xlsx')
    cell = book['Synthetic Data']['A2']
    assert cell.value == '=1+1'
    assert cell.data_type == 's'
    book.close()


def test_lock_and_company_change_are_blocked(tmp_path):
    c = client(tmp_path)
    run(tmp_path)
    lock = c / '.automation' / 'run.lock'
    lock.mkdir()
    assert 'locked' in run(tmp_path)['clients'][0]['error']
    lock.rmdir()
    config = read_json(c / 'client.json')
    config['company_id'] = 'different-company'
    write_json(c / 'client.json',config)
    assert 'History belongs' in run(tmp_path)['clients'][0]['error']


def test_cli_from_other_working_directory(tmp_path):
    c = client(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice())
    main = Path(__file__).resolve().parents[1] / 'main.py'
    result = subprocess.run([sys.executable,str(main),'--root',str(tmp_path),'run'],cwd='/',capture_output=True,text=True)
    assert result.returncode == 0
    assert 'Generated 1' in result.stdout
    assert read_json(c / 'reports' / 'validation.json')['counts']['new'] == 1


def test_same_numbers_new_financial_year_and_identical_numeric_format(tmp_path):
    c = client(tmp_path)
    original = invoice()
    pdf(c / INBOX / 'one.pdf',original)
    run(tmp_path)
    alternate = deepcopy(original)
    alternate['items'][0].update(rate='1000.00',quantity='2.0',value='2000.00',cgst_rate='9.00',sgst_rate='9.00')
    pdf(c / INBOX / 'two.pdf',alternate)
    report = run(tmp_path)['clients'][0]
    assert report['counts']['review'] == 0
    assert report['counts']['duplicate'] == 2
    alternate['date'] = '2027-04-01'
    pdf(c / INBOX / 'three.pdf',alternate)
    assert run(tmp_path)['clients'][0]['counts']['new'] == 1


def test_copied_company_id_blocks_both_clients(tmp_path):
    a = client(tmp_path)
    pdf(a / INBOX / 'one.pdf',invoice())
    fake_templates(tmp_path)
    run(tmp_path)
    shutil.copytree(a,tmp_path / 'Client 2')
    reports = run(tmp_path)['clients']
    assert len(reports) == 2
    assert all('share company_id' in r['error'] for r in reports)
    assert not read_json(a / 'imports' / 'manifest.json')['valid_for_import']


def test_changed_seller_cannot_reuse_history(tmp_path):
    from helpers import gstin
    c = client(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice())
    run(tmp_path)
    config = read_json(c / 'client.json')
    config['seller_gstin'] = gstin('29')
    write_json(c / 'client.json',config)
    assert 'different seller GSTIN' in run(tmp_path)['clients'][0]['error']


def test_current_snapshot_conflict_blocks_previously_pending_invoice(tmp_path):
    c = client(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice())
    fake_templates(tmp_path)
    run(tmp_path)
    (c / INBOX / 'one.pdf').unlink()
    changed = invoice()
    changed['buyer_name'] = 'Changed in Tally'
    config = read_json(c / 'client.json')
    write_snapshot(c / 'tally_snapshot.json',config,[{'invoice':changed}])
    report = run(tmp_path)['clients'][0]
    assert report['pending_validated'] == 0
    assert any('conflicts with existing Tally' in str(i) for i in report['issues'])
    assert not (c / 'imports' / 'Sales_Output.xlsx').exists()


def test_cli_template_inspection(tmp_path):
    folder,_ = fake_templates(tmp_path)
    result = inspect_workbook(folder / 'sales.xlsx')
    assert set(result['sheets']) == {'Synthetic Data','ReadMe'}
    assert result['sheets']['Synthetic Data']['sample_rows'][0] == ['Test Identifier','Test Value']


def test_rejected_and_pending_import_outcomes_retained(tmp_path):
    c = client(tmp_path)
    fake_templates(tmp_path)
    pdf(c / INBOX / 'one.pdf',invoice())
    run(tmp_path)
    manifest = read_json(c / 'imports' / 'manifest.json')
    result = {'company_id':manifest['company_id'],'generated_at':manifest['generated_at'],
              'tally_result_reference':'SYNTHETIC TEST ONLY',
              'masters':[{'master':m,'status':'rejected' if n == 0 else 'pending'} for n,m in enumerate(manifest['masters'])],
              'vouchers':[{'invoice_key':manifest['invoice_keys'][0],'status':'rejected'}]}
    path = c / 'import_results.json'
    write_json(path,result)
    assert confirm_import(c,path)[0]['status'] == 'rejected'
    db = sqlite3.connect(c / '.automation' / 'history.sqlite3')
    assert db.execute('SELECT COUNT(*) FROM import_results').fetchone()[0] == 1
    assert db.execute("SELECT COUNT(*) FROM events WHERE source LIKE 'master:%'").fetchone()[0] == len(manifest['masters'])
    db.close()
    assert run(tmp_path)['clients'][0]['pending_validated'] == 0


def test_pdfs_directly_in_client_folder_and_optional_inbox(tmp_path):
    c = client(tmp_path, 'client 1')
    pdf(c / 'direct.PDF',invoice('DIRECT'))
    pdf(c / INBOX / 'nested.pdf',invoice('NESTED'))
    pdf(c / 'reports' / 'ignore.pdf',invoice('NOT-AN-INPUT'))
    report = run(tmp_path)['clients'][0]
    assert report['counts']['new'] == 2
    records = read_json(c / 'reports' / 'validated_invoices.json')
    assert {r['invoice']['source'] for r in records} == {'direct.PDF',f'{INBOX}/nested.pdf'}
    assert run(tmp_path)['clients'][0]['counts']['duplicate'] == 2


def test_unconfigured_direct_pdf_folder_is_discovered(tmp_path):
    c = tmp_path / 'client 1'
    pdf(c / 'existing.pdf',invoice())
    report = run(tmp_path)['clients'][0]
    assert report['client'] == 'client 1'
    assert 'Missing client.json' in report['error']


def test_main_without_subcommand_and_relocated_root(tmp_path):
    source = Path(__file__).resolve().parents[1]
    root = tmp_path / 'dummy sales invoice'
    root.mkdir()
    shutil.copy2(source / 'main.py', root / 'main.py')
    shutil.copytree(source / 'Common_Rules',root / 'Common_Rules')
    c = client(root, 'client 1')
    pdf(c / 'one.pdf',invoice())
    result = subprocess.run([sys.executable,str(root / 'main.py')],cwd='/',capture_output=True,text=True)
    assert result.returncode == 0
    assert 'Generated 1' in result.stdout
    assert read_json(c / 'reports' / 'validation.json')['counts']['new'] == 1
